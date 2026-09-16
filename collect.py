"""매크로 트래커 로우데이터 수집 스크립트.

'지표목록' 시트를 설정표로 읽어 FRED에서 데이터를 받아 '로우데이터' 시트에 추가한다.
'로우데이터' 시트는 이미 전체 범위(A1:F5000)에 드롭다운·서식이 적용된 템플릿이 채워져
있으므로, openpyxl로 저장하면 x14 확장 데이터 유효성 검사(드롭다운)가 깨진다. 이를
피하기 위해 읽기는 openpyxl(read_only)로, 쓰기는 xlsx(zip) 내부의 시트 XML을 직접
문자열 치환하는 방식으로 처리한다.
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import warnings
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from xml.sax.saxutils import escape, unescape

import requests
from dotenv import load_dotenv

warnings.simplefilter("ignore", UserWarning)
import openpyxl  # noqa: E402  (경고 억제 이후 임포트)

XLSX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "매크로_트래커_로우데이터.xlsx")
SHEET_INDICATORS = "지표목록"
SHEET_RAWDATA = "로우데이터"
SHEET_LIST = "_목록"
RAWDATA_SHEET_XML = "xl/worksheets/sheet2.xml"  # 워크북 내 로우데이터 시트의 실제 파일명
INDICATORS_SHEET_XML = "xl/worksheets/sheet3.xml"  # 워크북 내 지표목록 시트의 실제 파일명
GUIDE_SHEET_XML = "xl/worksheets/sheet1.xml"  # 워크북 내 안내 시트의 실제 파일명
CALENDAR_SHEET_NAME = "발표일정"
CALENDAR_SHEET_XML = "xl/worksheets/sheet5.xml"  # 새로 추가하는 발표일정 시트
CALENDAR_HEADERS = ["예정일", "국가", "지표", "대상기간", "출처", "비고"]

# NO_RELEASE_DATE_FRED_SERIES/PREFIXES에 해당하는 (국가, 지표) — 기존에 이미 수집된
# 로우데이터 행의 발표일을 일괄로 비울 때 사용.
NO_RELEASE_DATE_INDICATOR_PAIRS = {
    ("일본", "실업률"),
    ("한국", "국채10년"),
    ("일본", "국채10년"),
    ("한국", "두바이유"),
    ("일본", "두바이유"),
}

GUIDE_NOTE_TEXT = "OECD·IMF 경유 FRED 시리즈는 FRED 반영일이 실제 발표일보다 늦어 발표일을 기록하지 않음"

# 로우데이터 B/C/F열 드롭다운이 참조하는 _목록 열: 국가(A)/지표(C)/주기(B)
LIST_COL_FOR_DROPDOWN = {"B": "A", "C": "C", "F": "B"}

FRED_BASE = "https://api.stlouisfed.org/fred/series/observations"
ECOS_BASE = "https://ecos.bok.or.kr/api"
ESTAT_BASE = "https://api.e-stat.go.jp/rest/3.0/app/json/getStatsData"
EUROSTAT_BASE = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data"
OBS_START = date(2021, 9, 1)

# OECD/IMF 경유 FRED 시리즈: FRED 반영일이 실제 발표일보다 늦어 ALFRED vintage를
# 신뢰할 수 없으므로 발표일을 조회하지 않고 항상 빈칸으로 저장한다.
NO_RELEASE_DATE_FRED_SERIES = {"LRUNTTTTJPM156S", "POILDUBUSDM"}
NO_RELEASE_DATE_FRED_PREFIXES = ("IRLTLT01",)


def _fred_release_date_supported(series_id):
    if series_id in NO_RELEASE_DATE_FRED_SERIES:
        return False
    return not series_id.startswith(NO_RELEASE_DATE_FRED_PREFIXES)


# 4단계 수집 대상 국가에 독일·프랑스 추가.
TARGET_COUNTRIES = {"미국", "한국", "일본", "독일", "프랑스"}

# 지표목록 시트 한국/일본 행의 '시리즈 ID' 칸(G열)에 채워 넣을 확정 코드.
# 행 번호는 지표목록 시트의 실제 위치(1~3단계 조사로 확정: 한국 30~34·36행, 일본 11~14·16행).
KOREA_SERIES_UPDATES = {
    30: "ECOS: 901Y009/0 (전년동월비 계산)",  # CPI 헤드라인
    31: "ECOS: 901Y010/QB (전년동월비 계산)",  # CPI 근원 (농산물 및 석유류 제외지수)
    32: "ECOS: 901Y027/I61BA+I28A (원계열)",  # 취업자수
    33: "ECOS: 901Y027/I61BC+I28B (계절조정)",  # 실업률
    34: "ECOS: 200Y108/10601 (전기비 계산)",  # 실질GDP 성장률
    36: "ECOS: 722Y001/0101000 (변경일만 기록)",  # 기준금리
}
JAPAN_SERIES_UPDATES = {
    11: "ESTAT: 0004052037/tab3/cat0001/area00000 (전년동월비, e-Stat 제공값)",  # CPI 헤드라인
    12: "ESTAT: 0004052037/tab3/cat0161/area00000 (전년동월비, e-Stat 제공값)",  # CPI 근원
    13: "FRED: LRUNTTTTJPM156S",  # 실업률 (계절조정, OECD 경유)
    14: "FRED: JPNRGDPEXP (전기비 계산)",  # 실질GDP 성장률
    16: (
        "BOJ 공표자료 (수동 확정: 2021-09-01=-0.10%, 2024-03-19=0.10%, "
        "2024-07-31=0.25%, 2025-01-24=0.50%, 2025-12-19=0.75%, 2026-06-16=1.00%)"
    ),  # 정책금리
}

# 일본은행 정책금리 변경 이력(BOJ 공표문 원문 대조로 확정, 2021-09~현재). 기준일=발표일=결정일.
JAPAN_POLICY_RATE_CHANGES = [
    (date(2021, 9, 1), -0.10),
    (date(2024, 3, 19), 0.10),
    (date(2024, 7, 31), 0.25),
    (date(2025, 1, 24), 0.50),
    (date(2025, 12, 19), 0.75),
    (date(2026, 6, 16), 1.00),
]

# HICP는 2026-01부터 ECOICOP 버전2(2025=100) 분류로 전환되며, prc_hicp_minr가 과거
# 시계열도 새 분류로 재계산해 제공한다(1996-01~). 4단계 사전 검증(2021-09~2025-12,
# 헤드라인·근원 모두 최대 diff 0.1%p, 문턱 0.3%p 이내)을 통과해 이 데이터셋 단일
# 소스로만 전체 기간을 수집한다(옛 prc_hicp_manr와 잇지 않음).
GERMANY_SERIES_UPDATES = {
    18: "EUROSTAT: prc_hicp_minr?geo=DE&unit=RCH_A&coicop18=TOTAL (ECOICOP v2, 전년동월비 직접 제공)",  # HICP 헤드라인
    19: "EUROSTAT: prc_hicp_minr?geo=DE&unit=RCH_A&coicop18=TOT_X_NRG_FOOD (ECOICOP v2, 에너지·식품·주류·담배 제외)",  # HICP 근원
    20: "EUROSTAT: une_rt_m?geo=DE&s_adj=SA&sex=T&age=TOTAL&unit=PC_ACT",  # 실업률
    21: "EUROSTAT: namq_10_gdp?geo=DE&na_item=B1GQ&unit=CLV_PCH_PRE&s_adj=SCA (전기비 직접 제공)",  # 실질GDP 성장률
    23: "EUROSTAT: irt_lt_mcby_m?geo=DE",  # 국채10년
}
FRANCE_SERIES_UPDATES = {
    24: "EUROSTAT: prc_hicp_minr?geo=FR&unit=RCH_A&coicop18=TOTAL (ECOICOP v2, 전년동월비 직접 제공)",  # HICP 헤드라인
    25: "EUROSTAT: prc_hicp_minr?geo=FR&unit=RCH_A&coicop18=TOT_X_NRG_FOOD (ECOICOP v2, 에너지·식품·주류·담배 제외)",  # HICP 근원
    26: "EUROSTAT: une_rt_m?geo=FR&s_adj=SA&sex=T&age=TOTAL&unit=PC_ACT",  # 실업률
    27: "EUROSTAT: namq_10_gdp?geo=FR&na_item=B1GQ&unit=CLV_PCH_PRE&s_adj=SCA (전기비 직접 제공)",  # 실질GDP 성장률
    29: "EUROSTAT: irt_lt_mcby_m?geo=FR",  # 국채10년
}

EXCEL_EPOCH = date(1899, 12, 30)


def to_excel_serial(d):
    return (d - EXCEL_EPOCH).days


def month_start(d):
    return date(d.year, d.month, 1)


def quarter_start(d):
    q_month = ((d.month - 1) // 3) * 3 + 1
    return date(d.year, q_month, 1)


def add_months(d, n):
    total = d.year * 12 + (d.month - 1) + n
    return date(total // 12, total % 12 + 1, 1)


# ---------------------------------------------------------------------------
# 지표목록 설정 로딩
# ---------------------------------------------------------------------------

def load_indicator_config(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[SHEET_INDICATORS]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    indicators = []
    for row in rows:
        if not row or row[0] is None:
            continue
        country, category, indicator, freq, unit, source, series_field, note = (
            list(row) + [None] * (8 - len(row))
        )[:8]
        indicators.append(
            {
                "country": country,
                "category": category,
                "indicator": indicator,
                "freq": freq,
                "unit": unit,
                "source": source,
                "series_field": series_field,
                "note": note,
            }
        )
    return indicators


FRED_FIELD_RE = re.compile(r"FRED:\s*([A-Za-z0-9_]+)")


def parse_fred_directive(series_field):
    if not series_field:
        return None
    m = FRED_FIELD_RE.search(series_field)
    if not m:
        return None
    return {
        "series_id": m.group(1),
        "is_pc1": "units=pc1" in series_field,
        "daily_to_monthly": "일간" in series_field and "월평균" in series_field,
        "qoq_from_level": "전기비" in series_field and "계산" in series_field,
    }


ECOS_FIELD_RE = re.compile(r"ECOS:\s*([A-Za-z0-9]+)/([A-Za-z0-9+]+)")


def parse_ecos_directive(series_field):
    if not series_field:
        return None
    m = ECOS_FIELD_RE.search(series_field)
    if not m:
        return None
    return {"stat_code": m.group(1), "item_codes": m.group(2).split("+")}


ESTAT_FIELD_RE = re.compile(r"ESTAT:\s*(\w+)/tab(\w+)/cat(\w+)/area(\w+)")


def parse_estat_directive(series_field):
    if not series_field:
        return None
    m = ESTAT_FIELD_RE.search(series_field)
    if not m:
        return None
    return {"stats_data_id": m.group(1), "tab": m.group(2), "cat01": m.group(3), "area": m.group(4)}


EUROSTAT_FIELD_RE = re.compile(r"EUROSTAT:\s*(\w+)\?(\S+)")


def parse_eurostat_directive(series_field):
    if not series_field:
        return None
    m = EUROSTAT_FIELD_RE.search(series_field)
    if not m:
        return None
    from urllib.parse import parse_qsl

    return {"dataset": m.group(1), "params": dict(parse_qsl(m.group(2)))}


BOJ_MANUAL_PREFIX = "BOJ 공표자료"


# ---------------------------------------------------------------------------
# FRED API
# ---------------------------------------------------------------------------

class FredClient:
    def __init__(self, api_key):
        self.api_key = api_key
        self.session = requests.Session()

    def _get(self, params):
        params = dict(params, api_key=self.api_key, file_type="json")
        resp = self.session.get(FRED_BASE, params=params, timeout=30)
        time.sleep(0.1)
        return resp

    def latest_observations(self, series_id, start, end, units="lin"):
        """최신(수정 반영) 값. {date: value_str} 반환, 결측('.')은 제외."""
        resp = self._get(
            {
                "series_id": series_id,
                "observation_start": start.isoformat(),
                "observation_end": end.isoformat(),
                "units": units,
            }
        )
        resp.raise_for_status()
        obs = resp.json().get("observations", [])
        return {o["date"]: o["value"] for o in obs if o["value"] != "."}

    def first_vintage_dates(self, series_id, dates):
        """주어진 관측일들(ISO 문자열 집합)에 대해 ALFRED 첫 vintage(최초 공개) 날짜를 구한다.

        전체 기간을 한 번에 조회(output_type=3)하되, vintage 개수 제한(2000개)에
        걸리면 관측일 하나씩 좁은 실시간 구간으로 재조회한다. 반환값도 ISO 문자열.
        """
        if not dates:
            return {}
        start, end = min(dates), max(dates)
        result = self._first_vintage_wide(series_id, start, end)
        if result is not None:
            return {d: v for d, v in result.items() if d in dates}

        # 넓은 구간 조회 실패 -> 관측일별 개별 조회로 폴백
        result = {}
        for d in sorted(dates):
            v = self._first_vintage_single(series_id, d)
            if v:
                result[d] = v
        return result

    def _first_vintage_wide(self, series_id, start, end):
        resp = self._get(
            {
                "series_id": series_id,
                "observation_start": start,
                "observation_end": end,
                "output_type": 3,
                "realtime_start": "1776-07-04",
                "realtime_end": "9999-12-31",
            }
        )
        if resp.status_code != 200:
            return None
        obs = resp.json().get("observations", [])
        out = {}
        for o in obs:
            vintages = [k for k in o.keys() if k != "date"]
            if not vintages:
                continue
            first_key = min(vintages, key=lambda k: k.rsplit("_", 1)[-1])
            out[o["date"]] = first_key.rsplit("_", 1)[-1]  # YYYYMMDD
        return {d: _yyyymmdd_to_iso(v) for d, v in out.items()}

    def _first_vintage_single(self, series_id, obs_date_str, window_days=60):
        obs_date = datetime.strptime(obs_date_str, "%Y-%m-%d").date()
        rt_start = obs_date - timedelta(days=window_days)
        rt_end = obs_date + timedelta(days=window_days)
        resp = self._get(
            {
                "series_id": series_id,
                "observation_start": obs_date_str,
                "observation_end": obs_date_str,
                "output_type": 3,
                "realtime_start": rt_start.isoformat(),
                "realtime_end": rt_end.isoformat(),
            }
        )
        if resp.status_code != 200:
            return None
        obs = resp.json().get("observations", [])
        if not obs:
            return None
        vintages = [k for k in obs[0].keys() if k != "date"]
        if not vintages:
            return None
        first_key = min(vintages, key=lambda k: k.rsplit("_", 1)[-1])
        return _yyyymmdd_to_iso(first_key.rsplit("_", 1)[-1])


def _yyyymmdd_to_iso(s):
    return f"{s[0:4]}-{s[4:6]}-{s[6:8]}"


# ---------------------------------------------------------------------------
# ECOS API
# ---------------------------------------------------------------------------

def _ecos_date(d, cycle):
    if cycle == "D":
        return f"{d.year:04d}{d.month:02d}{d.day:02d}"
    if cycle == "M":
        return f"{d.year:04d}{d.month:02d}"
    if cycle == "Q":
        return f"{d.year:04d}Q{(d.month - 1) // 3 + 1}"
    raise ValueError(f"지원하지 않는 ECOS 주기: {cycle}")


def _ecos_parse_time(t, cycle):
    if cycle == "D":
        return date(int(t[0:4]), int(t[4:6]), int(t[6:8]))
    if cycle == "M":
        return date(int(t[0:4]), int(t[4:6]), 1)
    if cycle == "Q":
        y, q = int(t[0:4]), int(t[5])
        return date(y, (q - 1) * 3 + 1, 1)
    raise ValueError(f"지원하지 않는 ECOS 주기: {cycle}")


class EcosClient:
    def __init__(self, api_key):
        self.api_key = api_key
        self.session = requests.Session()

    def observations(self, stat_code, cycle, start, end, item_codes, count=6000):
        """{date: float} 반환. ECOS는 vintage/발표일 정보를 제공하지 않는다."""
        parts = "/".join(item_codes)
        url = (
            f"{ECOS_BASE}/StatisticSearch/{self.api_key}/json/kr/1/{count}/"
            f"{stat_code}/{cycle}/{_ecos_date(start, cycle)}/{_ecos_date(end, cycle)}/{parts}"
        )
        resp = self.session.get(url, timeout=30)
        time.sleep(0.1)
        resp.raise_for_status()
        payload = resp.json()
        if "StatisticSearch" not in payload:
            err = payload.get("RESULT", payload)
            raise RuntimeError(f"ECOS API 오류: {err}")
        rows = payload["StatisticSearch"].get("row", [])
        out = {}
        for row in rows:
            try:
                v = float(row["DATA_VALUE"])
            except (TypeError, ValueError):
                continue
            out[_ecos_parse_time(row["TIME"], cycle)] = v
        return out


# ---------------------------------------------------------------------------
# e-Stat API
# ---------------------------------------------------------------------------

def _estat_time_code(d):
    return f"{d.year:04d}00{d.month:02d}{d.month:02d}"


class EstatClient:
    def __init__(self, app_id):
        self.app_id = app_id
        self.session = requests.Session()

    def observations(self, stats_data_id, tab, cat01, area, start, end):
        """{date: float} 반환(월간). e-Stat은 vintage/발표일 정보를 제공하지 않는다."""
        params = {
            "appId": self.app_id,
            "statsDataId": stats_data_id,
            "cdTab": tab,
            "cdCat01": cat01,
            "cdArea": area,
            "cdTimeFrom": _estat_time_code(start),
            "cdTimeTo": _estat_time_code(end),
        }
        resp = self.session.get(ESTAT_BASE, params=params, timeout=30)
        time.sleep(0.1)
        resp.raise_for_status()
        payload = resp.json()
        stat_data = payload.get("GET_STATS_DATA", {}).get("STATISTICAL_DATA")
        if stat_data is None:
            err = payload.get("GET_STATS_DATA", {}).get("RESULT", payload)
            raise RuntimeError(f"e-Stat API 오류: {err}")
        values = stat_data.get("DATA_INF", {}).get("VALUE", [])
        if isinstance(values, dict):
            values = [values]
        out = {}
        for v in values:
            t = v["@time"]
            d = date(int(t[0:4]), int(t[6:8]), 1)
            try:
                out[d] = float(v["$"])
            except (TypeError, ValueError):
                continue
        return out


# ---------------------------------------------------------------------------
# Eurostat API
# ---------------------------------------------------------------------------

def _eurostat_parse_time(t, cycle):
    if cycle == "Q":
        y, q = t.split("-Q")
        return date(int(y), (int(q) - 1) * 3 + 1, 1)
    y, m = t.split("-")
    return date(int(y), int(m), 1)


class EurostatClient:
    def __init__(self, max_retries=3):
        self.session = requests.Session()
        self.max_retries = max_retries

    def _get(self, dataset, params):
        url = f"{EUROSTAT_BASE}/{dataset}"
        last_exc = None
        for attempt in range(self.max_retries):
            try:
                resp = self.session.get(url, params={**params, "format": "JSON", "lang": "EN"}, timeout=30)
                resp.raise_for_status()
                return resp
            except requests.RequestException as e:  # noqa: PERF203
                last_exc = e
                time.sleep(1 + attempt)
        raise last_exc

    def observations(self, dataset, params, cycle):
        """{date: float} 반환. Eurostat JSON-stat 응답에는 vintage/발표일 정보가 없다.
        API의 startPeriod/sinceTimePeriod 필터가 불안정해 전체 시계열을 받은 뒤
        클라이언트에서 기간을 자른다."""
        resp = self._get(dataset, params)
        payload = resp.json()
        if "dimension" not in payload or "time" not in payload["dimension"]:
            err = payload.get("error", payload)
            raise RuntimeError(f"Eurostat API 오류: {err}")
        time_index = payload["dimension"]["time"]["category"]["index"]
        pos_to_time = {v: k for k, v in time_index.items()}
        out = {}
        for k, v in payload.get("value", {}).items():
            t = pos_to_time[int(k)]
            try:
                out[_eurostat_parse_time(t, cycle)] = float(v)
            except (TypeError, ValueError):
                continue
        return out


# ---------------------------------------------------------------------------
# 지표별 수집 로직
# ---------------------------------------------------------------------------

def collect_indicator(clients, cfg, today, start):
    """반환: (rows, error_message_or_None, held_back)

    rows: [{date, value, release_date_or_None}]
    held_back: [{series_id, date, value, vintage}] — 반복값 의심으로 저장하지 않은 최신월
    clients: {"fred": FredClient, "ecos": EcosClient, "estat": EstatClient, "eurostat": EurostatClient}
    start: 이 지표를 조회할 하한 날짜(증분 수집 시 재확인 구간 시작일, 신규 지표는 OBS_START)
    """
    fred_directive = parse_fred_directive(cfg["series_field"])
    if fred_directive is not None:
        return _collect_fred(clients["fred"], fred_directive, cfg["freq"], today, start)

    ecos_directive = parse_ecos_directive(cfg["series_field"])
    if ecos_directive is not None:
        return _collect_ecos(clients["ecos"], ecos_directive, cfg, today, start)

    estat_directive = parse_estat_directive(cfg["series_field"])
    if estat_directive is not None:
        return _collect_estat(clients["estat"], estat_directive, today, start)

    eurostat_directive = parse_eurostat_directive(cfg["series_field"])
    if eurostat_directive is not None:
        return _collect_eurostat(clients["eurostat"], eurostat_directive, cfg, today, start)

    if cfg["series_field"] and cfg["series_field"].startswith(BOJ_MANUAL_PREFIX):
        return _collect_boj_policy_rate(start), None, []

    return None, "지원하지 않는 시리즈 출처 (시리즈 ID: %r)" % cfg["series_field"], []


def _collect_fred(client, directive, freq, today, start):
    series_id = directive["series_id"]
    try:
        if freq == "수시":
            return _collect_adhoc(client, series_id, today, start), None, []
        if directive["daily_to_monthly"]:
            return _collect_daily_to_monthly(client, series_id, today, start), None, []
        if freq == "분기" and directive["qoq_from_level"]:
            return _collect_fred_qoq(client, series_id, today, start), None, []
        units = "pc1" if directive["is_pc1"] else "lin"
        rows, held_back = _collect_direct(client, series_id, units, today, start)
        return rows, None, held_back
    except requests.HTTPError as e:
        return None, f"FRED API 오류: {e}", []
    except Exception as e:  # noqa: BLE001
        return None, f"수집 실패: {e}", []


def _collect_ecos(client, directive, cfg, today, start):
    stat_code = directive["stat_code"]
    item_codes = directive["item_codes"]
    freq = cfg["freq"]
    category = cfg["category"]
    try:
        if freq == "수시":
            return _collect_ecos_adhoc(client, stat_code, item_codes, today, start), None, []
        if category == "물가":
            return _collect_ecos_yoy(client, stat_code, item_codes, today, start), None, []
        if category == "성장":
            return _collect_ecos_qoq(client, stat_code, item_codes, today, start), None, []
        return _collect_ecos_level(client, stat_code, item_codes, today, start), None, []
    except requests.HTTPError as e:
        return None, f"ECOS API 오류: {e}", []
    except Exception as e:  # noqa: BLE001
        return None, f"수집 실패: {e}", []


def _collect_direct(client, series_id, units, today, start):
    """월간(또는 pc1 변환) FRED 값을 그대로 저장. 최신 달 값이 직전 달과 완전히 같고
    ALFRED 게시일(vintage)도 같으면 최신 달은 반복값(placeholder)으로 의심해 저장하지
    않고 held_back에 기록한다."""
    values = client.latest_observations(series_id, start, today, units=units)
    if not values:
        return [], []

    supports_release = _fred_release_date_supported(series_id)
    dates_sorted = sorted(values.keys())

    if supports_release:
        vintage_map = client.first_vintage_dates(series_id, set(values.keys()))
    elif len(dates_sorted) >= 2:
        vintage_map = client.first_vintage_dates(series_id, set(dates_sorted[-2:]))
    else:
        vintage_map = {}

    skip = None
    held_back = []
    if len(dates_sorted) >= 2:
        latest, prev = dates_sorted[-1], dates_sorted[-2]
        latest_vintage = vintage_map.get(latest)
        if values[latest] == values[prev] and latest_vintage and latest_vintage == vintage_map.get(prev):
            skip = latest
            held_back.append(
                {
                    "series_id": series_id,
                    "date": datetime.strptime(latest, "%Y-%m-%d").date(),
                    "value": float(values[latest]),
                    "vintage": latest_vintage,
                    "suppressed": True,
                }
            )

    rows = []
    for d_str, v_str in sorted(values.items()):
        if d_str == skip:
            continue
        release_date = _parse_iso_or_none(vintage_map.get(d_str)) if supports_release else None
        rows.append(
            {
                "date": datetime.strptime(d_str, "%Y-%m-%d").date(),
                "value": round(float(v_str), 4),
                "release_date": release_date,
            }
        )
    return rows, held_back


def _collect_daily_to_monthly(client, series_id, today, start):
    daily = client.latest_observations(series_id, start, today, units="lin")
    if not daily:
        return []
    current_month = month_start(today)
    buckets = defaultdict(list)
    for d_str, v_str in daily.items():
        d = datetime.strptime(d_str, "%Y-%m-%d").date()
        ms = month_start(d)
        if ms >= current_month:
            continue  # 진행 중인 달(불완전 평균)은 제외
        buckets[ms].append(float(v_str))

    rows = []
    for ms in sorted(buckets.keys()):
        avg = sum(buckets[ms]) / len(buckets[ms])
        rows.append({"date": ms, "value": round(avg, 2), "release_date": None})
    return rows


def _collect_fred_qoq(client, series_id, today, start):
    """분기 레벨 시리즈를 받아 전기비 %로 변환(한국 GDP와 동일 방식). 직접 계산한
    파생값이라 ALFRED vintage가 의미 없으므로 발표일은 항상 빈칸."""
    target_start = quarter_start(start)
    lookback_start = add_months(target_start, -3)
    values = client.latest_observations(series_id, lookback_start, today, units="lin")
    ordered = sorted(values.items())  # (date_str, value_str)
    rows = []
    prev_v = None
    for d_str, v_str in ordered:
        d = datetime.strptime(d_str, "%Y-%m-%d").date()
        v = float(v_str)
        if prev_v is not None and d >= target_start:
            qoq = (v / prev_v - 1) * 100
            rows.append({"date": d, "value": round(qoq, 2), "release_date": None})
        prev_v = v
    return rows


def _collect_adhoc(client, series_id, today, start):
    """수시(정책금리) FRED 시리즈: 값이 바뀐 날짜만 기록. start 이전 이력도 함께 받아
    start 시점에 우연히 '변경일'로 오인되지 않도록 한 뒤, start 이후 구간만 낸다."""
    lookback_start = min(start - timedelta(days=400), OBS_START)
    daily = client.latest_observations(series_id, lookback_start, today, units="lin")
    if not daily:
        return []
    ordered = sorted(daily.items())  # (date_str, value_str)

    segments = []
    prev_val = None
    for d_str, v_str in ordered:
        v = float(v_str)
        if prev_val is None or v != prev_val:
            segments.append((d_str, v))
        prev_val = v

    segments = [(d_str, v) for d_str, v in segments if datetime.strptime(d_str, "%Y-%m-%d").date() >= start]

    seg_dates = {d for d, _ in segments}
    release_map = client.first_vintage_dates(series_id, seg_dates)

    rows = []
    for d_str, v in segments:
        rows.append(
            {
                "date": datetime.strptime(d_str, "%Y-%m-%d").date(),
                "value": round(v, 2),
                "release_date": _parse_iso_or_none(release_map.get(d_str)),
            }
        )
    return rows


def _parse_iso_or_none(s):
    if not s:
        return None
    return datetime.strptime(s, "%Y-%m-%d").date()


def _collect_ecos_yoy(client, stat_code, item_codes, today, start):
    """월간 지수를 받아 전년동월비 %로 변환. ECOS는 발표일을 주지 않으므로 항상 빈칸."""
    lookback_start = add_months(start, -12)
    values = client.observations(stat_code, "M", lookback_start, today, item_codes)
    rows = []
    for d in sorted(values):
        if d < start:
            continue
        prev = values.get(add_months(d, -12))
        if prev is None or prev == 0:
            continue
        yoy = (values[d] / prev - 1) * 100
        rows.append({"date": d, "value": round(yoy, 2), "release_date": None})
    return rows


def _collect_ecos_qoq(client, stat_code, item_codes, today, start):
    """분기 레벨을 받아 전기비 %로 변환. start가 속한 분기부터 포함(FRED 분기 시리즈와
    동일한 정렬 방식). ECOS는 발표일을 주지 않으므로 항상 빈칸."""
    target_start = quarter_start(start)
    lookback_start = add_months(target_start, -3)
    values = client.observations(stat_code, "Q", lookback_start, today, item_codes)
    ordered = sorted(values.items())
    rows = []
    prev_v = None
    for d, v in ordered:
        if prev_v is not None and d >= target_start:
            qoq = (v / prev_v - 1) * 100
            rows.append({"date": d, "value": round(qoq, 2), "release_date": None})
        prev_v = v
    return rows


def _collect_ecos_level(client, stat_code, item_codes, today, start):
    """레벨 값을 그대로 저장(취업자수/실업률 등). ECOS는 발표일을 주지 않으므로 항상 빈칸."""
    values = client.observations(stat_code, "M", start, today, item_codes)
    rows = []
    for d in sorted(values):
        if d < start:
            continue
        rows.append({"date": d, "value": round(values[d], 2), "release_date": None})
    return rows


def _collect_ecos_adhoc(client, stat_code, item_codes, today, start):
    """기준금리: 값이 바뀐 날짜만 기록. 기준일=발표일=금통위 결정일(변경 시작일)로 동일하게
    저장. start 이전의 변경 이력도 함께 받아, start가 우연히 '변경일'로 오인되지 않도록 한다."""
    lookback_start = min(date(start.year - 2, start.month, start.day), OBS_START)
    daily = client.observations(stat_code, "D", lookback_start, today, item_codes)
    ordered = sorted(daily.items())

    segments = []
    prev_val = None
    for d, v in ordered:
        if prev_val is None or v != prev_val:
            segments.append((d, v))
        prev_val = v

    rows = []
    for d, v in segments:
        if d < start:
            continue
        rows.append({"date": d, "value": round(v, 2), "release_date": d})
    return rows


def _collect_estat(client, directive, today, start):
    try:
        rows = _collect_estat_direct(
            client, directive["stats_data_id"], directive["tab"], directive["cat01"], directive["area"], today, start
        )
        return rows, None, []
    except requests.HTTPError as e:
        return None, f"e-Stat API 오류: {e}", []
    except Exception as e:  # noqa: BLE001
        return None, f"수집 실패: {e}", []


def _collect_estat_direct(client, stats_data_id, tab, cat01, area, today, start):
    """e-Stat이 이미 계산해 주는 값(전년동월비 등)을 그대로 저장. e-Stat은 발표일을
    제공하지 않으므로 항상 빈칸."""
    values = client.observations(stats_data_id, tab, cat01, area, start, today)
    rows = []
    for d in sorted(values):
        if d < start:
            continue
        rows.append({"date": d, "value": round(values[d], 2), "release_date": None})
    return rows


def _collect_eurostat(client, directive, cfg, today, start):
    dataset = directive["dataset"]
    params = directive["params"]
    cycle = "Q" if cfg["freq"] == "분기" else "M"
    try:
        rows, held_back = _collect_eurostat_direct(client, dataset, params, cycle, today, start)
        return rows, None, held_back
    except requests.HTTPError as e:
        return None, f"Eurostat API 오류: {e}", []
    except Exception as e:  # noqa: BLE001
        return None, f"수집 실패: {e}", []


def _collect_eurostat_direct(client, dataset, params, cycle, today, start):
    """Eurostat이 이미 계산해 주는 값(전년동월비·전기비 등)을 그대로 저장. Eurostat
    JSON-stat 응답에는 vintage/발표일이 없으므로 항상 빈칸. 같은 이유로 FRED처럼
    게시일 대조를 통한 반복값 확정 판정은 불가능하다 — 최신월이 직전월과 완전히
    같으면 저장은 하되(정상적인 flat 데이터일 수 있으므로 임의로 버리지 않음)
    held_back에 '검증 불가 반복값'으로만 표시해 보고서에서 확인할 수 있게 한다."""
    values = client.observations(dataset, params, cycle)
    # 분기 시리즈는 start가 속한 분기부터 포함(FRED/ECOS 분기 시리즈와 동일한 정렬 방식).
    # 월간은 start 그대로.
    start_bound = quarter_start(start) if cycle == "Q" else start
    filtered = {d: v for d, v in values.items() if start_bound <= d <= today}
    dates_sorted = sorted(filtered.keys())

    held_back = []
    if cycle == "M" and len(dates_sorted) >= 2:
        latest, prev = dates_sorted[-1], dates_sorted[-2]
        if filtered[latest] == filtered[prev]:
            held_back.append(
                {
                    "series_id": f"{dataset}?{params.get('geo', '')}",
                    "date": latest,
                    "value": filtered[latest],
                    "vintage": None,
                    "suppressed": False,
                }
            )

    rows = [{"date": d, "value": round(filtered[d], 2), "release_date": None} for d in dates_sorted]
    return rows, held_back


def _collect_boj_policy_rate(start):
    """일본은행 정책금리: BOJ 공표문을 대조해 수동으로 확정한 변경 이력(JAPAN_POLICY_RATE_CHANGES).
    기준일=발표일=금융정책결정회의일. 새 결정이 있으면 JAPAN_POLICY_RATE_CHANGES를 직접 갱신해야 한다."""
    rows = []
    for d, v in JAPAN_POLICY_RATE_CHANGES:
        if d < start:
            continue
        rows.append({"date": d, "value": v, "release_date": d})
    return rows


# ---------------------------------------------------------------------------
# 로우데이터 시트 읽기 (dedup, 다음 행 위치) / 쓰기 (XML 직접 치환)
# ---------------------------------------------------------------------------

def read_existing_rawdata(xlsx_path):
    """반환: (existing_keys, next_row, values_by_indicator)

    values_by_indicator: {(country, indicator): {date: value}} — 증분 수집의 시작일
    계산과 수정 감지(값 대조)에 쓰인다.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[SHEET_RAWDATA]
    existing_keys = set()
    values_by_indicator = defaultdict(dict)
    next_row = 2
    for row in ws.iter_rows(min_row=2, max_row=5000, values_only=True):
        base_date, country, indicator, value = row[0], row[1], row[2], row[3]
        if base_date is None:
            break
        if hasattr(base_date, "date"):
            base_date = base_date.date()
        existing_keys.add((base_date, country, indicator))
        values_by_indicator[(country, indicator)][base_date] = value
        next_row += 1
    wb.close()
    return existing_keys, next_row, dict(values_by_indicator)


def compute_indicator_start(freq, existing_dates):
    """증분 수집 시작일을 정한다. 기존 데이터가 없으면 OBS_START(전체 백필).
    월간/분기는 최근 3개(월/분기)를 재확인 구간으로 포함해 되돌아간다(수정 감지용).
    수시(정책금리)는 재확인 구간 개념이 없어 마지막 기준일 그대로 반환 — 각 수시
    수집 함수가 내부적으로 더 넓게 재조회해 변경 시점을 정확히 재판별한다."""
    if not existing_dates:
        return OBS_START
    last_date = max(existing_dates)
    if freq == "분기":
        return add_months(last_date, -6)  # 최근 3개 분기(당월 포함) 재확인
    if freq == "수시":
        return last_date
    return add_months(last_date, -2)  # 최근 3개월(당월 포함) 재확인


def _build_row_xml(row_num, base_date, country, indicator, value, release_date, freq):
    def date_cell(col, style, d):
        if d is None:
            return f'<c r="{col}{row_num}" s="{style}"/>'
        return f'<c r="{col}{row_num}" s="{style}"><v>{to_excel_serial(d)}</v></c>'

    def text_cell(col, style, text):
        return f'<c r="{col}{row_num}" s="{style}" t="inlineStr"><is><t>{escape(str(text))}</t></is></c>'

    def num_cell(col, style, num):
        return f'<c r="{col}{row_num}" s="{style}"><v>{num}</v></c>'

    a = date_cell("A", 6, base_date)
    b = text_cell("B", 7, country)
    c = text_cell("C", 7, indicator)
    d = num_cell("D", 7, value)
    e = date_cell("E", 6, release_date)
    f = text_cell("F", 7, freq)
    return f'<row r="{row_num}" spans="1:6">{a}{b}{c}{d}{e}{f}</row>'


def _empty_row_template(row_num):
    return (
        f'<row r="{row_num}" spans="1:6">'
        f'<c r="A{row_num}" s="6"/><c r="B{row_num}" s="7"/><c r="C{row_num}" s="7"/>'
        f'<c r="D{row_num}" s="7"/><c r="E{row_num}" s="6"/><c r="F{row_num}" s="7"/>'
        f"</row>"
    )


def _replace_sheet_xml(xlsx_path, sheet_path, new_xml):
    """xlsx(zip) 안의 sheet_path 항목만 new_xml로 교체하고 나머지는 그대로 둔다."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        names = zin.namelist()
        contents = {n: zin.read(n) for n in names}

    contents[sheet_path] = new_xml.encode("utf-8")

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(xlsx_path))
    os.close(fd)
    try:
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for n in names:
                zout.writestr(n, contents[n])
        shutil.move(tmp_path, xlsx_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


def write_new_rows(xlsx_path, new_records, start_row):
    """new_records: [(base_date, country, indicator, value, release_date, freq)]"""
    if not new_records:
        return

    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(RAWDATA_SHEET_XML).decode("utf-8")

    for i, rec in enumerate(new_records):
        row_num = start_row + i
        if row_num > 5000:
            raise RuntimeError("로우데이터 시트의 여유 행(5000행)을 초과했습니다.")
        old = _empty_row_template(row_num)
        if old not in sheet_xml:
            raise RuntimeError(f"{row_num}행이 예상한 빈 템플릿 형태가 아닙니다. 수동 확인이 필요합니다.")
        new = _build_row_xml(row_num, *rec)
        sheet_xml = sheet_xml.replace(old, new, 1)

    _replace_sheet_xml(xlsx_path, RAWDATA_SHEET_XML, sheet_xml)


_VALIDATION_BLOCK_RE = re.compile(r"<x14:dataValidation\b.*?</x14:dataValidation>", re.S)
_VALIDATION_F_RE = re.compile(r"(<xm:f>)(.*?)(</xm:f>)")
_VALIDATION_SQREF_RE = re.compile(r"<xm:sqref>([A-Z]+)\d+:[A-Z]+\d+</xm:sqref>")


def refresh_dropdown_ranges(xlsx_path):
    """로우데이터 B/C/F열 드롭다운(x14 데이터 유효성 검사)을 _목록 시트의 실제
    데이터 범위에 맞춰 다시 계산한다. _목록 시트 자체나 sqref(행 범위)는 건드리지
    않고, 각 드롭다운이 참조하는 '_목록!$X$1:$X$N' 범위의 N만 갱신한다."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    list_rows = list(wb[SHEET_LIST].iter_rows(values_only=True))
    wb.close()

    def col_count(col_letter):
        idx = ord(col_letter) - ord("A")
        n = 0
        for row in list_rows:
            if idx < len(row) and row[idx] is not None:
                n += 1
            else:
                break
        return n

    counts = {col: col_count(col) for col in set(LIST_COL_FOR_DROPDOWN.values())}
    for col, n in counts.items():
        if n < 1:
            raise RuntimeError(f"_목록 {col}열에서 값을 찾지 못했습니다.")

    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(RAWDATA_SHEET_XML).decode("utf-8")

    changed = False

    def _replace_block(m):
        nonlocal changed
        block = m.group(0)
        sq = _VALIDATION_SQREF_RE.search(block)
        if not sq:
            return block
        dropdown_col = sq.group(1)
        list_col = LIST_COL_FOR_DROPDOWN.get(dropdown_col)
        if not list_col:
            return block
        new_range = f"{SHEET_LIST}!${list_col}$1:${list_col}${counts[list_col]}"

        def _sub_f(fm):
            nonlocal changed
            if fm.group(2) != new_range:
                changed = True
            return fm.group(1) + new_range + fm.group(3)

        return _VALIDATION_F_RE.sub(_sub_f, block, count=1)

    new_sheet_xml = _VALIDATION_BLOCK_RE.sub(_replace_block, sheet_xml)

    if not changed:
        return False

    _replace_sheet_xml(xlsx_path, RAWDATA_SHEET_XML, new_sheet_xml)
    return True


def update_indicator_series_ids(xlsx_path, updates):
    """지표목록 시트 특정 행들의 '시리즈 ID'(G열)에 확정된 코드를 적는다. updates에
    정의된 행만 건드리고, 이미 갱신된 행은 건너뛴다. 행마다 기존 placeholder 텍스트가
    달라도(공유 문자열/스타일 무관) 정규식으로 G열 셀 하나만 정확히 찾아 통째로 교체한다."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(INDICATORS_SHEET_XML).decode("utf-8")

    changed = False
    for row_num, new_text in updates.items():
        new_cell = f'<c r="G{row_num}" s="8" t="inlineStr"><is><t>{escape(new_text)}</t></is></c>'
        if new_cell in sheet_xml:
            continue  # 이미 갱신됨
        cell_re = re.compile(r'<c r="G%d"(?:[^>]*?/>|[^>]*>.*?</c>)' % row_num)
        new_sheet_xml, n = cell_re.subn(new_cell, sheet_xml, count=1)
        if n != 1:
            raise RuntimeError(f"지표목록 G{row_num} 셀을 찾지 못했습니다. 수동 확인이 필요합니다.")
        sheet_xml = new_sheet_xml
        changed = True

    if not changed:
        return False

    _replace_sheet_xml(xlsx_path, INDICATORS_SHEET_XML, sheet_xml)
    return True


_ROW_BLOCK_RE = re.compile(r'<row r="(\d+)" spans="1:6">(.*?)</row>')


def _row_cell_text(row_body, col, row_num):
    m = re.search(r'<c r="%s%s"[^>]*t="inlineStr"><is><t>(.*?)</t></is></c>' % (col, row_num), row_body)
    return unescape(m.group(1)) if m else None


def clear_release_dates(xlsx_path, pairs):
    """로우데이터 시트에서 (국가, 지표)가 pairs에 속하는 기존 행들의 발표일(E열)을
    비운다. 이미 비어 있는 행은 건드리지 않는다."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(RAWDATA_SHEET_XML).decode("utf-8")

    changed = False

    def _process_row(m):
        nonlocal changed
        row_num, body = m.group(1), m.group(2)
        country = _row_cell_text(body, "B", row_num)
        indicator = _row_cell_text(body, "C", row_num)
        if (country, indicator) not in pairs:
            return m.group(0)
        blank_cell = f'<c r="E{row_num}" s="6"/>'
        e_cell_re = re.compile(r'<c r="E%s"[^/>]*/>|<c r="E%s"[^>]*>.*?</c>' % (row_num, row_num))
        e_match = e_cell_re.search(body)
        if not e_match or e_match.group(0) == blank_cell:
            return m.group(0)  # 발표일 셀이 없거나 이미 빈 상태
        new_body = e_cell_re.sub(blank_cell, body, count=1)
        changed = True
        return f'<row r="{row_num}" spans="1:6">{new_body}</row>'

    new_sheet_xml = _ROW_BLOCK_RE.sub(_process_row, sheet_xml)

    if not changed:
        return False

    _replace_sheet_xml(xlsx_path, RAWDATA_SHEET_XML, new_sheet_xml)
    return True


def delete_rawdata_rows(xlsx_path, keys_to_delete):
    """로우데이터 시트에서 (기준일, 국가, 지표)가 keys_to_delete에 속하는 행을 제거하고
    뒤따르는 행들을 한 칸씩 앞으로 당긴다(중간에 빈 행이 남지 않도록). 데이터 오류로
    잘못 수집된 값을 정정할 때 사용하는 일회성 유틸리티."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[SHEET_RAWDATA]
    kept, removed = [], []
    for row in ws.iter_rows(min_row=2, max_row=5000, values_only=True):
        if row[0] is None:
            break
        base_date = row[0].date() if hasattr(row[0], "date") else row[0]
        release_date = row[4].date() if hasattr(row[4], "date") else row[4]
        rec = (base_date, row[1], row[2], row[3], release_date, row[5])
        (removed if (base_date, row[1], row[2]) in keys_to_delete else kept).append(rec)
    wb.close()

    if not removed:
        return []

    last_row = 1 + len(kept) + len(removed)

    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(RAWDATA_SHEET_XML).decode("utf-8")

    def _rebuild(m):
        row_num = int(m.group(1))
        if row_num < 2 or row_num > last_row:
            return m.group(0)
        idx = row_num - 2
        if idx < len(kept):
            return _build_row_xml(row_num, *kept[idx])
        return _empty_row_template(row_num)

    new_sheet_xml = _ROW_BLOCK_RE.sub(_rebuild, sheet_xml)
    _replace_sheet_xml(xlsx_path, RAWDATA_SHEET_XML, new_sheet_xml)
    return removed


def insert_guide_note(xlsx_path):
    """'안내' 시트의 '수정치' 항목(입력 규칙 목록) 바로 아래에 FRED 발표일 정책
    한 줄을 새 행으로 삽입한다. 이미 삽입돼 있으면 아무것도 하지 않는다(멱등)."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(GUIDE_SHEET_XML).decode("utf-8")
        shared = zin.read("xl/sharedStrings.xml").decode("utf-8")

    if GUIDE_NOTE_TEXT in sheet_xml:
        return False

    m = re.search(r'<row r="19"[^>]*>.*?</row>', sheet_xml)
    a19 = re.search(r'<c r="A19"[^>]*t="s"><v>(\d+)</v></c>', m.group(0)) if m else None
    a19_text = None
    if a19:
        strings = re.findall(r"<si><t[^>]*>(.*?)</t></si>", shared, re.S)
        idx = int(a19.group(1))
        if idx < len(strings):
            a19_text = unescape(strings[idx])
    if a19_text != "수정치":
        raise RuntimeError("'안내' 시트 19행이 예상한 '수정치' 항목이 아닙니다. 수동 확인이 필요합니다.")

    # 20~22행을 21~23행으로 한 칸씩 밀어낸다(뒤에서부터 치환해야 충돌하지 않는다).
    for old_row in (22, 21, 20):
        new_row = old_row + 1
        old_block = re.search(r'<row r="%d"[^>]*>.*?</row>' % old_row, sheet_xml).group(0)
        shifted = old_block.replace(f'r="{old_row}"', f'r="{new_row}"', 1)
        shifted = re.sub(r'r="([A-Z]+)%d"' % old_row, r'r="\g<1>%d"' % new_row, shifted)
        sheet_xml = sheet_xml.replace(old_block, shifted, 1)

    new_row_xml = (
        '<row r="20" spans="1:2">'
        f'<c r="A20" s="3" t="inlineStr"><is><t>{escape("FRED 발표일")}</t></is></c>'
        f'<c r="B20" s="2" t="inlineStr"><is><t>{escape(GUIDE_NOTE_TEXT)}</t></is></c>'
        "</row>"
    )
    anchor = re.search(r'<row r="19"[^>]*>.*?</row>', sheet_xml).group(0)
    sheet_xml = sheet_xml.replace(anchor, anchor + new_row_xml, 1)

    sheet_xml = sheet_xml.replace('<dimension ref="A1:B22"/>', '<dimension ref="A1:B23"/>', 1)

    _replace_sheet_xml(xlsx_path, GUIDE_SHEET_XML, sheet_xml)
    return True


# ---------------------------------------------------------------------------
# 발표일정 시트 (신규 시트 생성 + 행 추가/조회)
# ---------------------------------------------------------------------------

def create_calendar_sheet(xlsx_path):
    """워크북에 '발표일정' 시트를 새로 추가한다(헤더 행만). 이미 있으면 아무것도
    하지 않는다(멱등). 기존 시트(안내/로우데이터/지표목록/_목록)는 전혀 건드리지 않는다."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        names = zin.namelist()
        if CALENDAR_SHEET_XML in names:
            return False
        contents = {n: zin.read(n) for n in names}

    ct = contents["[Content_Types].xml"].decode("utf-8")
    override = (
        f'<Override PartName="/{CALENDAR_SHEET_XML[3:]}" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    )
    ct = ct.replace("</Types>", override + "</Types>")
    contents["[Content_Types].xml"] = ct.encode("utf-8")

    rels = contents["xl/_rels/workbook.xml.rels"].decode("utf-8")
    existing_rids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels)]
    new_rid = max(existing_rids) + 1
    new_rel = (
        f'<Relationship Id="rId{new_rid}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        f'Target="worksheets/{CALENDAR_SHEET_XML.rsplit("/", 1)[-1]}"/>'
    )
    rels = rels.replace("</Relationships>", new_rel + "</Relationships>")
    contents["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    wbxml = contents["xl/workbook.xml"].decode("utf-8")
    existing_sheet_ids = [int(m) for m in re.findall(r'sheetId="(\d+)"', wbxml)]
    new_sheet_id = max(existing_sheet_ids) + 1
    new_sheet_tag = f'<sheet name="{CALENDAR_SHEET_NAME}" sheetId="{new_sheet_id}" r:id="rId{new_rid}"/>'
    wbxml = wbxml.replace("</sheets>", new_sheet_tag + "</sheets>")
    contents["xl/workbook.xml"] = wbxml.encode("utf-8")

    header_cells = "".join(
        f'<c r="{col}1" s="5" t="inlineStr"><is><t>{escape(h)}</t></is></c>'
        for col, h in zip("ABCDEF", CALENDAR_HEADERS)
    )
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<dimension ref="A1:F1"/>'
        '<sheetViews><sheetView workbookViewId="0">'
        '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        "</sheetView></sheetViews>"
        '<sheetFormatPr baseColWidth="10" defaultColWidth="8.83203125" defaultRowHeight="17"/>'
        "<cols>"
        '<col min="1" max="1" width="12" customWidth="1"/>'
        '<col min="2" max="2" width="10" customWidth="1"/>'
        '<col min="3" max="3" width="22" customWidth="1"/>'
        '<col min="4" max="4" width="12" customWidth="1"/>'
        '<col min="5" max="5" width="16" customWidth="1"/>'
        '<col min="6" max="6" width="34" customWidth="1"/>'
        "</cols>"
        f'<sheetData><row r="1" spans="1:6" ht="22" customHeight="1">{header_cells}</row></sheetData>'
        "</worksheet>"
    )
    contents[CALENDAR_SHEET_XML] = sheet_xml.encode("utf-8")

    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(xlsx_path))
    os.close(fd)
    try:
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for n in names:
                zout.writestr(n, contents[n])
            zout.writestr(CALENDAR_SHEET_XML, contents[CALENDAR_SHEET_XML])
        shutil.move(tmp_path, xlsx_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise
    return True


def _row_cell_date(row_body, col, row_num):
    m = re.search(r'<c r="%s%s"[^>]*><v>(\d+)</v></c>' % (col, row_num), row_body)
    if not m:
        return None
    return EXCEL_EPOCH + timedelta(days=int(m.group(1)))


def append_calendar_rows(xlsx_path, rows):
    """rows: [(예정일: date, 국가, 지표, 대상기간, 출처, 비고)]. 이미 있는
    (예정일,국가,지표,대상기간) 조합은 건너뛴다(멱등, 중복 방지)."""
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        sheet_xml = zin.read(CALENDAR_SHEET_XML).decode("utf-8")

    existing = set()
    last_row = 1
    for m in re.finditer(r'<row r="(\d+)" spans="1:6"[^>]*>(.*?)</row>', sheet_xml):
        row_num = int(m.group(1))
        last_row = max(last_row, row_num)
        if row_num == 1:
            continue
        body = m.group(2)
        d = _row_cell_date(body, "A", row_num)
        country = _row_cell_text(body, "B", row_num)
        indicator = _row_cell_text(body, "C", row_num)
        period = _row_cell_text(body, "D", row_num)
        existing.add((d, country, indicator, period))

    new_rows_xml = []
    added = 0
    row_num = last_row
    for 예정일, 국가, 지표, 대상기간, 출처, 비고 in rows:
        key = (예정일, 국가, 지표, 대상기간)
        if key in existing:
            continue
        row_num += 1
        cells = (
            f'<c r="A{row_num}" s="6"><v>{to_excel_serial(예정일)}</v></c>'
            f'<c r="B{row_num}" s="7" t="inlineStr"><is><t>{escape(국가)}</t></is></c>'
            f'<c r="C{row_num}" s="7" t="inlineStr"><is><t>{escape(지표)}</t></is></c>'
            f'<c r="D{row_num}" s="7" t="inlineStr"><is><t>{escape(대상기간)}</t></is></c>'
            f'<c r="E{row_num}" s="7" t="inlineStr"><is><t>{escape(출처)}</t></is></c>'
            f'<c r="F{row_num}" s="7" t="inlineStr"><is><t>{escape(비고 or "")}</t></is></c>'
        )
        new_rows_xml.append(f'<row r="{row_num}" spans="1:6">{cells}</row>')
        existing.add(key)
        added += 1

    if not new_rows_xml:
        return 0

    sheet_xml = sheet_xml.replace("</sheetData>", "".join(new_rows_xml) + "</sheetData>")
    sheet_xml = re.sub(r'<dimension ref="A1:F\d+"/>', f'<dimension ref="A1:F{row_num}"/>', sheet_xml)

    _replace_sheet_xml(xlsx_path, CALENDAR_SHEET_XML, sheet_xml)
    return added


def read_calendar(xlsx_path):
    """{(국가, 지표, 대상기간): 예정일} 반환. 시트가 없으면 빈 dict."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if CALENDAR_SHEET_NAME not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[CALENDAR_SHEET_NAME]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        d = row[0].date() if hasattr(row[0], "date") else row[0]
        out[(row[1], row[2], row[3])] = d
    wb.close()
    return out


def target_period(d, freq):
    if freq == "분기":
        return f"{d.year}-Q{(d.month - 1) // 3 + 1}"
    return f"{d.year}-{d.month:02d}"


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

# 발표일정을 조회할 수 있는 카테고리(통계 발표). 유가/금리(시장가격·정책금리)는
# 스스로 발표일이 정해지거나(정책금리) 발표 개념이 없어(시장가격) 조회하지 않는다.
CALENDAR_LOOKUP_CATEGORIES = {"물가", "고용", "성장"}

LOG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logs")


def main():
    log_lines = []

    def emit(msg=""):
        print(msg)
        log_lines.append(msg)

    load_dotenv(dotenv_path=os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))
    fred_key = os.environ.get("FRED_API_KEY")
    if not fred_key:
        print("오류: .env 파일에 FRED_API_KEY가 설정되어 있지 않습니다.", file=sys.stderr)
        sys.exit(1)
    ecos_key = os.environ.get("ECOS_API_KEY")
    if not ecos_key:
        print("오류: .env 파일에 ECOS_API_KEY가 설정되어 있지 않습니다.", file=sys.stderr)
        sys.exit(1)
    estat_key = os.environ.get("ESTAT_APP_ID")
    if not estat_key:
        print("오류: .env 파일에 ESTAT_APP_ID가 설정되어 있지 않습니다.", file=sys.stderr)
        sys.exit(1)

    today = date.today()
    clients = {
        "fred": FredClient(fred_key),
        "ecos": EcosClient(ecos_key),
        "estat": EstatClient(estat_key),
        "eurostat": EurostatClient(),
    }

    emit(f"=== 매크로 트래커 업데이트 실행: {today.isoformat()} ===\n")

    series_ids_updated = update_indicator_series_ids(
        XLSX_PATH,
        {**KOREA_SERIES_UPDATES, **JAPAN_SERIES_UPDATES, **GERMANY_SERIES_UPDATES, **FRANCE_SERIES_UPDATES},
    )
    emit(f"=== 지표목록 한국·일본·독일·프랑스 행 시리즈 ID {'갱신함' if series_ids_updated else '이미 최신 상태'} ===\n")

    release_dates_cleared = clear_release_dates(XLSX_PATH, NO_RELEASE_DATE_INDICATOR_PAIRS)
    emit(f"=== OECD·IMF 경유 시리즈 기존 발표일 {'비움' if release_dates_cleared else '이미 빈 상태'} ===")

    guide_note_added = insert_guide_note(XLSX_PATH)
    emit(f"=== '안내' 시트 FRED 발표일 정책 안내문 {'추가함' if guide_note_added else '이미 있음'} ===")

    calendar_created = create_calendar_sheet(XLSX_PATH)
    emit(f"=== '발표일정' 시트 {'새로 생성함' if calendar_created else '이미 있음'} ===\n")

    calendar = read_calendar(XLSX_PATH)

    all_indicators = load_indicator_config(XLSX_PATH)
    indicators = [i for i in all_indicators if i["country"] in TARGET_COUNTRIES]
    emit(f"'{SHEET_INDICATORS}' 시트에서 {len(indicators)}개 지표(대상 국가: {', '.join(sorted(TARGET_COUNTRIES))})를 확인했습니다.\n")

    existing_keys, next_row, values_by_indicator = read_existing_rawdata(XLSX_PATH)
    emit(f"기존 '{SHEET_RAWDATA}' 행: {len(existing_keys)}개, 다음 입력 행: {next_row}\n")

    to_write = []
    summary = []  # (indicator_label, new_rows_for_print)
    failed = []
    all_held_back = []  # (indicator_label, held_back_dict)
    revisions = []  # (indicator_label, date, old_value, new_value)

    for cfg in indicators:
        label = f"{cfg['country']} / {cfg['indicator']}"
        existing_map = values_by_indicator.get((cfg["country"], cfg["indicator"]), {})
        start = compute_indicator_start(cfg["freq"], existing_map)

        rows, err, held_back = collect_indicator(clients, cfg, today, start)
        for hb in held_back:
            all_held_back.append((label, hb))
        if err:
            failed.append(f"{label}: {err}")
            continue
        if not rows:
            continue  # 재확인 구간에 신규/기존 데이터가 전혀 없을 수 있음(정상)

        new_rows = []
        for r in rows:
            key = (r["date"], cfg["country"], cfg["indicator"])
            if r["date"] in existing_map:
                old_value = existing_map[r["date"]]
                if old_value is not None and abs(float(old_value) - float(r["value"])) > 1e-9:
                    revisions.append((label, r["date"], old_value, r["value"]))
                continue  # 기존 행은 절대 덮어쓰지 않음
            if key in existing_keys:
                continue

            release_date = r["release_date"]
            if release_date is None and cfg["category"] in CALENDAR_LOOKUP_CATEGORIES:
                period = target_period(r["date"], cfg["freq"])
                release_date = calendar.get((cfg["country"], cfg["indicator"], period))

            to_write.append((r["date"], cfg["country"], cfg["indicator"], r["value"], release_date, cfg["freq"]))
            existing_keys.add(key)
            new_rows.append({**r, "release_date": release_date})

        if new_rows:
            summary.append((label, new_rows))

    if to_write:
        write_new_rows(XLSX_PATH, to_write, next_row)

    dropdown_updated = refresh_dropdown_ranges(XLSX_PATH)
    emit(f"=== 수집 결과: {len(to_write)}개 행 추가 ===")
    emit(f"=== B/C/F열 드롭다운(_목록 기준) {'갱신함' if dropdown_updated else '이미 최신 상태'} ===\n")

    if summary:
        emit("=== 지표별 추가 행 ===")
        for label, new_rows in summary:
            emit(f"[{label}] 신규 {len(new_rows)}행:")
            for r in new_rows:
                rel = r["release_date"].isoformat() if r["release_date"] else "(발표일 없음)"
                emit(f"  {r['date'].isoformat()} | 값={r['value']} | 발표일={rel}")
        emit("")
    else:
        emit("=== 신규로 추가된 행이 없습니다 (모든 지표가 최신 상태) ===\n")

    if revisions:
        emit("=== 수정 감지 (원자료 값이 바뀌었지만 기존 행은 덮어쓰지 않음) ===")
        for label, d, old_v, new_v in revisions:
            emit(f"  - {label} {d.isoformat()}: 기존={old_v} -> 원자료 최신={new_v}")
        emit("")

    if failed:
        emit("=== 수집하지 못한 지표 ===")
        for f in failed:
            emit(f"  - {f}")
    else:
        emit("모든 대상 지표를 확인했습니다.")

    suppressed = [(l, hb) for l, hb in all_held_back if hb.get("suppressed")]
    noted = [(l, hb) for l, hb in all_held_back if not hb.get("suppressed")]
    if suppressed:
        emit("\n=== 반복값 의심으로 보류 (직전 달과 값·FRED 게시일이 모두 동일해 저장하지 않음) ===")
        for label, hb in suppressed:
            emit(f"  - {label} ({hb['series_id']}) {hb['date'].isoformat()}: 값={hb['value']}, 게시일={hb['vintage']}")
    if noted:
        emit("\n=== 반복값 발견(검증 불가, 저장은 함 — Eurostat은 게시일 정보가 없어 확정 판정 불가) ===")
        for label, hb in noted:
            emit(f"  - {label} ({hb['series_id']}) {hb['date'].isoformat()}: 값={hb['value']}")

    os.makedirs(LOG_DIR, exist_ok=True)
    log_path = os.path.join(LOG_DIR, f"update_{today.strftime('%Y%m%d')}.txt")
    with open(log_path, "w", encoding="utf-8") as f:
        f.write("\n".join(log_lines) + "\n")
    print(f"\n(실행 로그 저장: {log_path})")

    export_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "export.py")
    if os.path.exists(export_path):
        print("\n=== export.py 실행 ===")
        result = subprocess.run([sys.executable, export_path])
        if result.returncode != 0:
            print("경고: export.py 실행 중 오류가 발생했습니다.", file=sys.stderr)


if __name__ == "__main__":
    main()
