"""매크로 트래커 엑셀 -> 홈페이지용 JSON/엑셀 내보내기.

'매크로_트래커_로우데이터.xlsx'를 읽기 전용으로 열어 site/data/data.json(로우데이터
전체 + 지표목록의 단위·출처)과 site/data/calendar.json(발표일정), 그리고
site/data/downloads/ 아래 국가별·전체·회귀분석 엑셀 다운로드 파일을 생성한다.
엑셀 원본은 전혀 수정하지 않는다(읽기만 함). collect.py 실행이 끝나면 자동으로
이어서 실행되지만, 단독으로도 실행할 수 있다: `python3 export.py`.
"""

import json
import os
import warnings
from datetime import date, datetime

warnings.simplefilter("ignore", UserWarning)
import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

XLSX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "매크로_트래커_로우데이터.xlsx")
SITE_DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "site", "data")
DATA_JSON_PATH = os.path.join(SITE_DATA_DIR, "data.json")
CALENDAR_JSON_PATH = os.path.join(SITE_DATA_DIR, "calendar.json")
ANALYSIS_JSON_PATH = os.path.join(SITE_DATA_DIR, "analysis.json")
DOWNLOADS_DIR = os.path.join(SITE_DATA_DIR, "downloads")

CATEGORY_ORDER = ["물가", "고용", "성장", "유가", "금리"]
COUNTRY_FILE = {
    "한국": "korea.xlsx",
    "미국": "usa.xlsx",
    "일본": "japan.xlsx",
    "독일": "germany.xlsx",
    "프랑스": "france.xlsx",
}
RAW_DATA_FILE = "raw_data.xlsx"
REGRESSION_FILE = "regression_analysis.xlsx"


def _to_iso(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    return v


def export_indicators(wb):
    ws = wb["지표목록"]
    indicators = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        country, category, indicator, freq, unit, source = (list(row) + [None] * 6)[:6]
        indicators.append(
            {
                "country": country,
                "category": category,
                "indicator": indicator,
                "freq": freq,
                "unit": unit,
                "source": source,
            }
        )
    return indicators


def export_observations(wb):
    ws = wb["로우데이터"]
    observations = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        base_date, country, indicator, value, release_date, freq = (list(row) + [None] * 6)[:6]
        observations.append(
            {
                "date": _to_iso(base_date),
                "country": country,
                "indicator": indicator,
                "value": value,
                "release_date": _to_iso(release_date),
                "freq": freq,
            }
        )
    return observations


def export_calendar(wb):
    if "발표일정" not in wb.sheetnames:
        return []
    ws = wb["발표일정"]
    events = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        scheduled_date, country, indicator, period, source, note = (list(row) + [None] * 6)[:6]
        events.append(
            {
                "date": _to_iso(scheduled_date),
                "country": country,
                "indicator": indicator,
                "period": period,
                "source": source,
                "note": note,
            }
        )
    events.sort(key=lambda e: e["date"] or "")
    return events


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 40)


def _write_rows(ws, header, rows):
    ws.append(header)
    for r in rows:
        ws.append(r)
    _autosize(ws)


def build_country_workbook(country, indicators, observations, out_path):
    """국가별 엑셀 1개: 시트 = 분류(물가/고용/성장/유가/금리), 열 = 기준일 | 지표별 값 | 발표일."""
    cat_indicators = {}
    for i in indicators:
        if i["country"] != country:
            continue
        cat_indicators.setdefault(i["category"], [])
        if i["indicator"] not in cat_indicators[i["category"]]:
            cat_indicators[i["category"]].append(i["indicator"])

    obs_lookup = {}  # (category, indicator) -> {date: (value, release_date)}
    for o in observations:
        if o["country"] != country:
            continue
        cat = next((c for c, inds in cat_indicators.items() if o["indicator"] in inds), None)
        if cat is None:
            continue
        obs_lookup.setdefault((cat, o["indicator"]), {})[o["date"]] = (o["value"], o["release_date"])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for cat in CATEGORY_ORDER:
        inds = cat_indicators.get(cat)
        if not inds:
            continue
        ws = wb.create_sheet(title=cat)
        all_dates = sorted({d for ind in inds for d in obs_lookup.get((cat, ind), {})})
        header = ["기준일"] + inds + ["발표일"]
        rows = []
        for d in all_dates:
            row = [d]
            release = None
            for ind in inds:
                v = obs_lookup.get((cat, ind), {}).get(d)
                row.append(v[0] if v else None)
                if v and v[1] and release is None:
                    release = v[1]
            row.append(release)
            rows.append(row)
        _write_rows(ws, header, rows)

    if not wb.sheetnames:
        ws = wb.create_sheet(title="데이터없음")
        ws.append(["이 국가에 대해 내보낼 지표가 없습니다."])
    wb.save(out_path)


def build_raw_data_workbook(observations, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "로우데이터"
    header = ["기준일", "국가", "지표", "값", "발표일", "주기"]
    rows = [[o["date"], o["country"], o["indicator"], o["value"], o["release_date"], o["freq"]] for o in observations]
    _write_rows(ws, header, rows)
    wb.save(out_path)


def build_regression_workbook(analysis, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    country_label = {"US": "미국", "JP": "일본", "DE": "독일", "FR": "프랑스"}

    market = analysis.get("market", {})
    ws1 = wb.create_sheet("단순회귀(시장금리)")
    rows = []
    for code, r in market.get("simple", {}).items():
        if r.get("insufficient"):
            continue
        rows.append(
            [country_label.get(code, code), r["beta"], r["se"], r["t"], r["p"], r["r2"], r["n"], r["hac_lags"]]
        )
    _write_rows(ws1, ["국가", "계수(β)", "표준오차(HAC)", "t값", "p값", "R²", "표본수", "HAC시차"], rows)

    multi = market.get("multi", {})
    ws2 = wb.create_sheet("다중회귀(시장금리)")
    coeffs = multi.get("coeffs", {})
    pvals = multi.get("p_values", {})
    vif = multi.get("vif", {})
    rows = []
    for var in multi.get("included", []):
        rows.append([country_label.get(var, var), coeffs.get(var), pvals.get(var), vif.get(var)])
    _write_rows(ws2, ["변수", "계수", "p값", "VIF"], rows)
    ws2.append([])
    ws2.append(["상수항", coeffs.get("const"), pvals.get("const"), None])
    ws2.append(["조정 R²", multi.get("adj_r2")])
    ws2.append(["R²", multi.get("r2")])
    ws2.append(["표본수", multi.get("n")])
    for code, info in multi.get("excluded", {}).items():
        ws2.append([f"제외: {country_label.get(code, code)}", info.get("reason"), "DE와 상관계수", info.get("corr_with_DE")])
    _autosize(ws2)

    ws3 = wb.create_sheet("동월vs1개월선행")
    rows = []
    for code, ll in market.get("lead_lag", {}).items():
        c, l1 = ll.get("contemporaneous", {}), ll.get("lead1", {})
        rows.append(
            [
                country_label.get(code, code),
                c.get("beta"), c.get("p"), c.get("r2"),
                l1.get("beta"), l1.get("p"), l1.get("r2"),
            ]
        )
    _write_rows(
        ws3,
        ["국가", "동월 β", "동월 p값", "동월 R²", "1개월선행 β", "1개월선행 p값", "1개월선행 R²"],
        rows,
    )

    policy = analysis.get("policy", {})
    ws4 = wb.create_sheet("정책금리 스프레드")
    spread = policy.get("spread_kr_us", {})
    dates = spread.get("dates", [])
    values = spread.get("KR_US", [])
    spread_jp = policy.get("spread_kr_jp", {})
    jp_map = dict(zip(spread_jp.get("dates", []), spread_jp.get("KR_JP", [])))
    rows = [[d, v, jp_map.get(d)] for d, v in zip(dates, values)]
    _write_rows(ws4, ["월말", "한미 정책금리차(%p)", "한일 정책금리차(%p)"], rows)

    ws5 = wb.create_sheet("교차상관(한국-미국)")
    cc = policy.get("cross_corr_kr_us", {})
    rows = list(zip(cc.get("lags", []), cc.get("corr", [])))
    _write_rows(ws5, ["시차(개월, 양수=미국 선행)", "상관계수"], rows)

    wb.save(out_path)


FX_DAILY_CSV_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "fx_daily.csv")
FX_DAILY_FILE = "fx_daily.xlsx"
FX_MONTHLY_FILE = "fx_monthly.xlsx"
FX_PAIR_ORDER = ["원/달러", "원/100엔", "원/유로", "달러지수(광의)", "엔/달러", "달러/유로"]


def build_fx_daily_workbook(out_path):
    import csv as _csv

    if not os.path.exists(FX_DAILY_CSV_PATH):
        return False
    with open(FX_DAILY_CSV_PATH, encoding="utf-8-sig", newline="") as f:
        rows_raw = list(_csv.DictReader(f))
    by_date = {}
    for r in rows_raw:
        by_date.setdefault(r["날짜"], {})[r["통화쌍"]] = float(r["값"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "일별"
    header = ["날짜"] + FX_PAIR_ORDER
    rows = []
    for d in sorted(by_date.keys()):
        row = [d] + [by_date[d].get(pair) for pair in FX_PAIR_ORDER]
        rows.append(row)
    _write_rows(ws, header, rows)
    wb.save(out_path)
    return True


def build_fx_monthly_workbook(fx_json, out_path):
    monthly = fx_json.get("monthly", {})
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "월평균"
    ws2 = wb.create_sheet("월말")
    for ws, key in ((ws1, "avg"), (ws2, "eom")):
        all_dates = sorted(set(d for pair in FX_PAIR_ORDER for d in monthly.get(pair, {}).get("dates", [])))
        header = ["연월"] + FX_PAIR_ORDER
        rows = []
        for d in all_dates:
            row = [d[:7]]
            for pair in FX_PAIR_ORDER:
                m = monthly.get(pair, {})
                dates = m.get("dates", [])
                vals = m.get(key, [])
                row.append(vals[dates.index(d)] if d in dates else None)
            rows.append(row)
        _write_rows(ws, header, rows)
    wb.save(out_path)


def export_downloads(indicators, observations):
    os.makedirs(DOWNLOADS_DIR, exist_ok=True)
    for country, fname in COUNTRY_FILE.items():
        build_country_workbook(country, indicators, observations, os.path.join(DOWNLOADS_DIR, fname))
    build_raw_data_workbook(observations, os.path.join(DOWNLOADS_DIR, RAW_DATA_FILE))

    if os.path.exists(ANALYSIS_JSON_PATH):
        with open(ANALYSIS_JSON_PATH, encoding="utf-8") as f:
            analysis = json.load(f)
        build_regression_workbook(analysis, os.path.join(DOWNLOADS_DIR, REGRESSION_FILE))
        regression_written = True
    else:
        regression_written = False

    fx_written = build_fx_daily_workbook(os.path.join(DOWNLOADS_DIR, FX_DAILY_FILE))
    fx_json_path = os.path.join(SITE_DATA_DIR, "fx.json")
    if fx_written and os.path.exists(fx_json_path):
        with open(fx_json_path, encoding="utf-8") as f:
            fx_json = json.load(f)
        build_fx_monthly_workbook(fx_json, os.path.join(DOWNLOADS_DIR, FX_MONTHLY_FILE))

    return regression_written


def main():
    if not os.path.exists(XLSX_PATH):
        raise SystemExit(f"오류: {XLSX_PATH} 파일을 찾을 수 없습니다.")

    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    indicators = export_indicators(wb)
    observations = export_observations(wb)
    calendar_events = export_calendar(wb)
    wb.close()

    os.makedirs(SITE_DATA_DIR, exist_ok=True)

    generated_at = datetime.now().isoformat(timespec="seconds")

    with open(DATA_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(
            {"generated_at": generated_at, "indicators": indicators, "observations": observations},
            f,
            ensure_ascii=False,
            indent=2,
        )

    with open(CALENDAR_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump({"generated_at": generated_at, "events": calendar_events}, f, ensure_ascii=False, indent=2)

    print(f"export.py: {DATA_JSON_PATH} 생성 (지표 {len(indicators)}개, 관측치 {len(observations)}개)")
    print(f"export.py: {CALENDAR_JSON_PATH} 생성 (발표일정 {len(calendar_events)}개)")

    regression_written = export_downloads(indicators, observations)
    print(f"export.py: {DOWNLOADS_DIR} 아래 국가별·전체 엑셀 생성" + ("(회귀분석 엑셀 포함)" if regression_written else "(analysis.json 없어 회귀분석 엑셀은 건너뜀)"))


if __name__ == "__main__":
    main()
